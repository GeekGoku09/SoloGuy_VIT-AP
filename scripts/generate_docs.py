import os
import json
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import letter, landscape
from reportlab.pdfgen import canvas

def create_architecture_pptx(path="docs/architecture.pptx"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    tf = txBox.text_frame
    tf.text = "Digital Twin Architecture (Prediabetes & T2D India Context)"
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True

    boxes = [
        ("1. Data Streams", "Static EHR + Dynamic CGM, Wearable, Meals"),
        ("2. Harmonisation", "5-Min Canonical Grid, Causal Forward-Fill"),
        ("3. Twin Core", "PatientState (Baseline, Gain, Decay, Uncertainty)"),
        ("4. Fusion Forecaster", "LightGBM + Quantile Trajectories + What-If"),
        ("5. Doctor Dashboard", "6-Panel Streamlit Dashboard & AGP")
    ]

    left = Inches(0.5)
    top = Inches(2.0)
    width = Inches(1.7)
    height = Inches(3.0)

    for idx, (title, desc) in enumerate(boxes):
        shape = slide.shapes.add_shape(1, left + idx * Inches(1.85), top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0, 114, 178) if idx < 4 else RGBColor(0, 158, 115)
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(14)
        p1.font.color.rgb = RGBColor(255, 255, 255)
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(240, 240, 240)

    prs.save(path)
    print(f"Generated {path}")

def create_architecture_pdf(path="docs/architecture.pdf"):
    c = canvas.Canvas(path, pagesize=landscape(letter))
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 540, "Digital Twin Architecture (Prediabetes & T2D India Context)")

    boxes = [
        ("1. Data Streams", "Static EHR + Dynamic CGM, Wearable"),
        ("2. Harmonisation", "5-Min Canonical Grid, Causal Fill"),
        ("3. Twin Core", "PatientState (Baseline, Gain, Uncertainty)"),
        ("4. Fusion Forecaster", "LightGBM + Quantile Bands"),
        ("5. Doctor Dashboard", "6-Panel Dashboard & AGP")
    ]

    for idx, (title, desc) in enumerate(boxes):
        x = 50 + idx * 140
        c.setFillColorRGB(0.0, 0.45, 0.70) if idx < 4 else c.setFillColorRGB(0.0, 0.62, 0.45)
        c.rect(x, 200, 130, 280, fill=1)
        c.setFillColorRGB(1, 1, 1)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x + 10, 450, title)
        c.setFont("Helvetica", 10)
        c.drawString(x + 10, 400, desc[:25])
        c.drawString(x + 10, 385, desc[25:])

    c.save()
    print(f"Generated {path}")

def create_presentation_pptx(path="docs/presentation.pptx"):
    prs = Presentation()
    
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "PoC Digital Twin for T2D & Prediabetes"
    slide1.placeholders[1].text = "Multi-Stream Fusion Forecaster & What-If Engine\nTEAM_NAME (2026)"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "The India Context & Healthcare Challenge"
    slide2.placeholders[1].text = "• 101 Million people with diabetes in India\n• 136 Million people with prediabetes (ICMR-INDIAB 2023)\n• Acute postprandial glucose spikes (>180 mg/dL) drive microvascular complications\n• Need for predictive forward warning & personalized intervention"

    prs.save(path)
    print(f"Generated {path}")

def create_presentation_pdf(path="docs/presentation.pdf"):
    c = canvas.Canvas(path, pagesize=landscape(letter))
    c.setFont("Helvetica-Bold", 24)
    c.drawString(100, 400, "PoC Digital Twin for T2D & Prediabetes")
    c.setFont("Helvetica", 16)
    c.drawString(100, 350, "Multi-Stream Fusion Forecaster & What-If Engine")
    c.drawString(100, 320, "TEAM_NAME (2026)")
    c.showPage()

    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 500, "The India Context & Healthcare Challenge")
    c.setFont("Helvetica", 14)
    c.drawString(50, 450, "• 101 Million people with diabetes in India")
    c.drawString(50, 420, "• 136 Million people with prediabetes (ICMR-INDIAB 2023)")
    c.drawString(50, 390, "• Acute postprandial glucose spikes (>180 mg/dL) drive microvascular complications")
    c.drawString(50, 360, "• Need for predictive forward warning & personalized intervention")
    c.save()
    print(f"Generated {path}")

if __name__ == "__main__":
    os.makedirs("docs", exist_ok=True)
    create_architecture_pptx()
    create_architecture_pdf()
    create_presentation_pptx()
    create_presentation_pdf()
